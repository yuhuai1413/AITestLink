from __future__ import annotations

import os
import uuid

from fastapi import HTTPException
from sqlalchemy import select

from app.config import settings
from app.models.file_asset import FileAsset
from app.models.project import Project
from app.models.requirement import Requirement
from app.models.status_log import StatusLog
from app.services.base import BaseService
from app.contracts.document import RequirementUpdate
from app.services.requirement_clarification import (
    CLARIFICATION_NOT_REQUIRED,
    CLARIFICATION_PENDING,
    clarification_answer_quality_issues,
    default_clarification_status,
    is_clarification_resolved,
)

ALLOWED_EXTENSIONS = {
    "docx", "doc", "pdf", "md", "txt", "json", "yaml", "yml", "csv",
    "xlsx", "xls", "png", "jpg", "jpeg",
}
MAX_FILE_SIZE = 50 * 1024 * 1024


class DocumentService(BaseService):
    """文档管理服务"""

    async def upload(self, file_content: bytes, filename: str, project_id: str) -> dict:
        if not filename or "." not in filename:
            from fastapi import HTTPException
            raise HTTPException(status_code=400, detail="文件必须包含扩展名")

        ext = filename.rsplit(".", 1)[-1].lower()
        if ext not in ALLOWED_EXTENSIONS:
            from fastapi import HTTPException
            raise HTTPException(
                status_code=400,
                detail=f"不支持的文件格式: .{ext}，支持: {', '.join(sorted(ALLOWED_EXTENSIONS))}",
            )

        if len(file_content) > MAX_FILE_SIZE:
            from fastapi import HTTPException
            raise HTTPException(status_code=400, detail="文件大小超过 50MB 限制")

        # 保存文件
        upload_dir = os.path.join(settings.UPLOAD_DIR, project_id)
        os.makedirs(upload_dir, exist_ok=True)

        file_id = str(uuid.uuid4())
        safe_filename = f"{file_id}.{ext}"
        filepath = os.path.join(upload_dir, safe_filename)

        with open(filepath, "wb") as f:
            f.write(file_content)

        file_asset = FileAsset(
            id=file_id,
            project_id=project_id,
            name=filename,
            file_type=ext,
            size=f"{len(file_content) / 1024:.1f} KB",
            storage_path=filepath,
        )
        self.db.add(file_asset)

        # 上传文件后，若项目测试状态为「待测试」则自动变为「测试中」
        proj_result = await self.db.execute(select(Project).where(Project.id == project_id))
        project = proj_result.scalar_one_or_none()
        if project and project.test_status == "待测试":
            project.test_status = "测试中"
            self.db.add(StatusLog(
                project_id=project_id,
                user_id=project.user_id,
                field_name="test_status",
                old_value="待测试",
                new_value="测试中",
                change_type="auto",
                reason="上传了输入资料，进入测试流程",
            ))

        await self.db.commit()
        await self.db.refresh(file_asset)
        return self._to_dict(file_asset)

    async def list_by_project(self, project_id: str) -> list[dict]:
        result = await self.db.execute(
            select(FileAsset).where(FileAsset.project_id == project_id).order_by(FileAsset.uploaded_at.desc())
        )
        return [self._to_dict(f) for f in result.scalars().all()]

    async def get_by_id(self, doc_id: str) -> dict | None:
        result = await self.db.execute(select(FileAsset).where(FileAsset.id == doc_id))
        file_asset = result.scalar_one_or_none()
        if not file_asset:
            return None
        return self._to_dict(file_asset)

    async def delete(self, doc_id: str) -> bool:
        from sqlalchemy import delete
        result = await self.db.execute(select(FileAsset).where(FileAsset.id == doc_id))
        file_asset = result.scalar_one_or_none()
        if not file_asset:
            return False

        project_id = file_asset.project_id

        # 删除物理文件
        if file_asset.storage_path and os.path.exists(file_asset.storage_path):
            os.remove(file_asset.storage_path)

        await self.db.delete(file_asset)

        # 级联删除关联数据：需求、测试点、测试用例、自动化脚本
        from app.models.requirement import Requirement
        from app.models.test_point import TestPoint
        from app.models.test_case import TestCase
        from app.models.automation_script import AutomationScript
        from app.models.execution_run import ExecutionRun

        await self.db.execute(delete(ExecutionRun).where(ExecutionRun.project_id == project_id))
        await self.db.execute(delete(Requirement).where(Requirement.project_id == project_id))
        await self.db.execute(delete(TestPoint).where(TestPoint.project_id == project_id))
        await self.db.execute(delete(TestCase).where(TestCase.project_id == project_id))
        await self.db.execute(delete(AutomationScript).where(AutomationScript.project_id == project_id))

        await self.db.commit()
        return True

    async def get_content(self, doc_id: str) -> str | None:
        """读取文档文本内容"""
        result = await self.db.execute(select(FileAsset).where(FileAsset.id == doc_id))
        file_asset = result.scalar_one_or_none()
        if not file_asset or not file_asset.storage_path:
            return None

        if not os.path.exists(file_asset.storage_path):
            return None

        ext = file_asset.file_type.lower()

        if ext in ("txt", "md", "json", "yaml", "yml", "csv"):
            with open(file_asset.storage_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()

        if ext in ("docx", "doc"):
            try:
                import docx
                doc = docx.Document(file_asset.storage_path)
                return "\n".join([p.text for p in doc.paragraphs])
            except ImportError:
                return "[docx 库未安装，无法解析 Word 文件]"

        if ext == "pdf":
            try:
                from PyPDF2 import PdfReader
                reader = PdfReader(file_asset.storage_path)
                texts = []
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        texts.append(text)
                return "\n".join(texts) if texts else "[PDF 文件无法提取文本内容]"
            except Exception:
                return "[PDF 解析失败]"

        if ext in ("xlsx", "xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_asset.storage_path, read_only=True)
                texts = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        texts.append(" | ".join([str(c) if c is not None else "" for c in row]))
                return "\n".join(texts)
            except ImportError:
                return "[openpyxl 库未安装，无法解析 Excel 文件]"

        return f"[不支持的文件类型: {ext}]"

    async def get_requirements(self, project_id: str) -> list[dict]:
        result = await self.db.execute(
            select(Requirement).where(Requirement.project_id == project_id).order_by(Requirement.created_at)
        )
        return [self._to_dict(r) for r in result.scalars().all()]

    async def update_requirement(self, req_id: str, data: RequirementUpdate) -> dict | None:
        result = await self.db.execute(select(Requirement).where(Requirement.id == req_id))
        req = result.scalar_one_or_none()
        if not req:
            return None

        update_data = data.model_dump(exclude_unset=True)
        field_map = {
            "rule": "rule",
            "question": "question",
            "confirmed": "confirmed",
            "clarificationStatus": "clarification_status",
            "clarification_status": "clarification_status",
            "clarificationAnswer": "clarification_answer",
            "clarification_answer": "clarification_answer",
            "reviewStatus": "review_status",
            "review_status": "review_status",
        }
        for schema_key, db_key in field_map.items():
            if schema_key in update_data:
                setattr(req, db_key, update_data[schema_key])

        # 已失效的需求不允许再次评审通过——数据失效后需重新解析需求，
        # 不能在旧数据上恢复评审状态。防止前端绕过或同步逻辑误改。
        from app.services.data_lineage_service import INVALID, REVIEW_INVALIDATED
        new_review = update_data.get("review_status")
        is_invalid = req.validity_status == INVALID or req.review_status == REVIEW_INVALIDATED
        if new_review == "已通过" and is_invalid:
            raise HTTPException(status_code=400, detail="该需求已失效，无法再次评审。请重新解析需求")

        req.clarification_status = default_clarification_status(
            req.question,
            False,
            req.clarification_answer,
        )
        req.confirmed = req.clarification_status in {"已确认", "无需确认"}

        if req.review_status == "已通过" and not is_clarification_resolved(req.question, req.clarification_status, req.clarification_answer):
            issues = clarification_answer_quality_issues(req.question, req.clarification_answer)
            detail = "该需求还有待确认问题未处理，请先填写确认结论"
            if issues and (req.clarification_answer or "").strip():
                detail = "确认结论仍不充分：" + "；".join(issues)
            raise HTTPException(status_code=400, detail=detail)

        req.updated_at = self._now()
        await self.db.commit()
        await self.db.refresh(req)
        return self._to_dict(req)

    async def delete_requirement(self, req_id: str) -> bool:
        from app.services.data_lineage_service import cascade_delete_requirement

        deleted = await cascade_delete_requirement(self.db, req_id)
        if not deleted:
            return False
        await self.db.commit()
        return True

    async def search_requirements(self, project_id: str, query: str) -> list[dict]:
        result = await self.db.execute(
            select(Requirement).where(
                Requirement.project_id == project_id,
                (Requirement.module.contains(query))
                | (Requirement.feature.contains(query))
                | (Requirement.rule.contains(query))
            )
        )
        return [self._to_dict(r) for r in result.scalars().all()]

    async def batch_confirm(self, req_ids: list[str], confirmed: bool) -> int:
        result = await self.db.execute(
            select(Requirement).where(Requirement.id.in_(req_ids))
        )
        count = 0
        for req in result.scalars().all():
            req.clarification_status = default_clarification_status(
                req.question,
                False,
                req.clarification_answer,
            )
            req.confirmed = req.clarification_status != CLARIFICATION_PENDING
            count += 1
        await self.db.commit()
        return count
